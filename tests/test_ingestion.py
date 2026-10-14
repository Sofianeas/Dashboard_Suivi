
from openpyxl import Workbook

from src.ingestion.google_sheets import load_from_xlsx


def test_load_from_xlsx_reads_tickets_and_reference_mappings(tmp_path):
    """Vérifie le chargement du classeur sans données opérationnelles."""
    workbook = Workbook()

    # Feuille des tickets : les en-têtes commencent à la ligne 5.
    tickets_sheet = workbook.active
    tickets_sheet.title = "Suivi Tickets"

    tickets_sheet.cell(row=5, column=1, value="Date du ticket")
    tickets_sheet.cell(row=5, column=2, value="N° Ticket")
    tickets_sheet.cell(row=5, column=3, value="Client")
    tickets_sheet.cell(row=5, column=4, value="Problématique")

    tickets_sheet.cell(row=6, column=1, value="23/06/2026")
    tickets_sheet.cell(row=6, column=2, value="P111111111")
    tickets_sheet.cell(row=6, column=3, value="CLIENT TEST")
    tickets_sheet.cell(row=6, column=4, value="Incident")

    # Feuille des référentiels clients et problématiques.
    lists_sheet = workbook.create_sheet("Listes")

    lists_sheet["H1"] = "Client source"
    lists_sheet["I1"] = "Client analytique"
    lists_sheet["H2"] = "CLIENT TEST"
    lists_sheet["I2"] = "CLIENT TEST"

    lists_sheet["L1"] = "Problème source"
    lists_sheet["M1"] = "Problème analytique"
    lists_sheet["L2"] = "Incident"
    lists_sheet["M2"] = "Incident"

    workbook_path = tmp_path / "classeur_synthetique.xlsx"
    workbook.save(workbook_path)
    workbook.close()

    tickets, clients, problems = load_from_xlsx(workbook_path)

    # Vérification du chargement des tickets.
    assert len(tickets) == 1
    assert tickets.loc[0, "N° Ticket"] == "P111111111"
    assert tickets.loc[0, "Client"] == "CLIENT TEST"
    assert tickets.loc[0, "Problématique"] == "Incident"

    # Vérification du mapping clients.
    assert list(clients.columns) == [
        "client_source_normalized",
        "client_analytic",
    ]
    assert len(clients) == 1
    assert clients.loc[0, "client_source_normalized"] == "CLIENT TEST"
    assert clients.loc[0, "client_analytic"] == "CLIENT TEST"

    # Vérification du mapping des problématiques.
    assert list(problems.columns) == [
        "problem_source",
        "problem_analytic",
    ]
    assert len(problems) == 1
    assert problems.loc[0, "problem_source"] == "Incident"
    assert problems.loc[0, "problem_analytic"] == "Incident"
