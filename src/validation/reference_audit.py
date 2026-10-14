"""Compare l'ETL V1 aux colonnes analytiques déjà présentes dans le classeur.
Ce rapport sert à distinguer les écarts du pipeline des particularités historiques du fichier.
"""
from __future__ import annotations
import openpyxl
import pandas as pd
from pathlib import Path
from ..ingestion.google_sheets import load_from_xlsx
from ..transformation.pipeline import transform_tickets


def audit_against_xlsx(path: str | Path) -> dict:
    path = Path(path)
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Suivi Tickets"]
    tickets, clients, problems = load_from_xlsx(path)
    result = transform_tickets(tickets, clients, problems)

    refs = []
    for row in range(6, ws.max_row + 1):
        ticket = ws.cell(row, 2).value
        if ticket:
            refs.append({
                "row": row,
                "ticket_id": str(ticket).strip(),
                "J": ws.cell(row, 10).value,
                "K": ws.cell(row, 11).value,
                "L": ws.cell(row, 12).value,
                "Z": ws.cell(row, 26).value,
            })
    ref = pd.DataFrame(refs)

    duration_mismatches = []
    for i, expected in enumerate(ref["J"]):
        actual = result.iloc[i]["resolution_minutes"]
        if pd.notna(expected) and abs(float(actual) - float(expected)) > 1e-6:
            duration_mismatches.append({"row": int(ref.iloc[i].row), "ticket_id": ref.iloc[i].ticket_id,
                                        "etl": float(actual), "sheet_J": float(expected)})

    def text_mismatches(col_result, col_ref):
        out=[]
        for i, expected in enumerate(ref[col_ref]):
            actual=result.iloc[i][col_result]
            a="" if pd.isna(actual) else str(actual)
            e="" if pd.isna(expected) else str(expected)
            if a != e:
                out.append({"row":int(ref.iloc[i].row),"ticket_id":ref.iloc[i].ticket_id,"etl":a,"sheet":e})
        return out

    return {
        "ticket_rows": len(result),
        "duration_mismatches_vs_J": duration_mismatches,
        "client_mismatches_vs_K": text_mismatches("client_analytic", "K"),
        "problem_mismatches_vs_L": text_mismatches("problem_analytic", "L"),
        "technician_mismatches_vs_Z": text_mismatches("technician_analytic", "Z"),
        "invalid_opened_at": int(result.opened_at.isna().sum()),
        "invalid_resolved_at": int(result.resolved_at.isna().sum()),
        "duplicate_ticket_ids_distinct": int(result.loc[result.ticket_id.duplicated(keep=False), "ticket_id"].nunique()),
        "exact_duplicate_rows": int(result.exact_duplicate.sum()),
    }


if __name__ == "__main__":
    import argparse, json
    p = argparse.ArgumentParser()
    p.add_argument("xlsx")
    p.add_argument("--output", default="data/processed/reference_audit.json")
    args=p.parse_args()
    report=audit_against_xlsx(args.xlsx)
    Path(args.output).parent.mkdir(parents=True, exist_ok=True)
    Path(args.output).write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
